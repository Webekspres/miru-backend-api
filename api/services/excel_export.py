"""
Excel export service for MIRU reports (Modul 16).

Generates .xlsx files for daily, weekly, monthly, waste, and evaluation reports
using openpyxl. All exports include institution header, report period, and
formatted data tables.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .reports import (
    daily_report,
    evaluation_report,
    monthly_report,
    waste_report,
    weekly_report,
)

# ── Style constants ────────────────────────────────────────
HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
HEADER_FONT = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
TITLE_FONT = Font(name='Calibri', bold=True, size=14)
SUBTITLE_FONT = Font(name='Calibri', bold=True, size=12)
NORMAL_FONT = Font(name='Calibri', size=11)
BOLD_FONT = Font(name='Calibri', bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)


def _style_header_row(ws, row: int, max_col: int) -> None:
    """Apply header styling to a row."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER


def _style_cell(ws, row: int, col: int, bold: bool = False) -> None:
    cell = ws.cell(row=row, column=col)
    cell.font = BOLD_FONT if bold else NORMAL_FONT
    cell.border = THIN_BORDER
    cell.alignment = Alignment(vertical='center', wrap_text=True)


def _write_kop(ws, judul: str, periode: str | None = None) -> int:
    """Write institution header and return current row number."""
    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value='MIRU Bank Sampah - Distrik Mimika Baru').font = TITLE_FONT
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=judul).font = SUBTITLE_FONT
    row += 1
    if periode:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1, value=f'Periode: {periode}').font = NORMAL_FONT
        row += 1
    row += 1  # blank row
    return row


def _write_section_title(ws, row: int, title: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=title).font = SUBTITLE_FONT
    return row + 1


def _auto_width(ws, max_col: int) -> None:
    """Auto-adjust column widths."""
    for col in range(1, max_col + 1):
        max_len = 0
        letter = get_column_letter(col)
        for cell in ws[letter]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 3, 40)


# ── Main export functions ─────────────────────────────────


def export_daily(tanggal: str | None) -> bytes:
    """Export daily report as .xlsx bytes."""
    data = daily_report(tanggal)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Laporan Harian'

    row = _write_kop(ws, 'LAPORAN HARIAN', data.get('tanggal', tanggal))

    # Summary
    headers = ['Jumlah Transaksi', 'Total Setoran', 'Total Penarikan', 'Total Sampah (kg)']
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(headers))
    row += 1

    values = [
        data['jumlah_transaksi'],
        data['total_setoran'],
        data['total_penarikan'],
        data['total_sampah_kg'],
    ]
    for i, v in enumerate(values, 1):
        ws.cell(row=row, column=i, value=v)
        _style_cell(ws, row, i)
    row += 2

    # Tonase per jenis
    row = _write_section_title(ws, row, 'Tonase per Jenis')
    tonase_headers = ['No', 'Kategori', 'Total Berat (kg)', 'Total Nilai']
    for i, h in enumerate(tonase_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(tonase_headers))
    row += 1

    for idx, item in enumerate(data.get('tonase_per_jenis', []), 1):
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=item['nama'])
        ws.cell(row=row, column=3, value=item['total_berat_kg'])
        ws.cell(row=row, column=4, value=item['total_nilai'])
        for c in range(1, 5):
            _style_cell(ws, row, c)
        row += 1

    # Footer
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=f'Dicetak: {datetime.now().strftime("%d/%m/%Y %H:%M")}').font = NORMAL_FONT

    _auto_width(ws, 6)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_weekly(minggu: str | None, tahun: str | None) -> bytes:
    """Export weekly report as .xlsx bytes."""
    data = weekly_report(minggu, tahun)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Laporan Mingguan'

    periode = data.get('periode', {})
    periode_str = f'{periode.get("mulai", "")} s.d. {periode.get("selesai", "")}'
    row = _write_kop(ws, f'LAPORAN MINGGUAN (Minggu ke-{data["minggu"]}, {data["tahun"]})', periode_str)

    # Summary
    headers = ['Jumlah Transaksi', 'Total Setoran', 'Total Penarikan', 'Total Sampah (kg)', 'Nasabah Baru']
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(headers))
    row += 1

    values = [
        data['jumlah_transaksi'],
        data['total_setoran'],
        data['total_penarikan'],
        data['total_sampah_kg'],
        data['nasabah_baru'],
    ]
    for i, v in enumerate(values, 1):
        ws.cell(row=row, column=i, value=v)
        _style_cell(ws, row, i)
    row += 2

    # Tonase per jenis
    row = _write_section_title(ws, row, 'Tonase per Jenis')
    tonase_headers = ['No', 'Kategori', 'Total Berat (kg)', 'Total Nilai']
    for i, h in enumerate(tonase_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(tonase_headers))
    row += 1

    for idx, item in enumerate(data.get('tonase_per_jenis', []), 1):
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=item['nama'])
        ws.cell(row=row, column=3, value=item['total_berat_kg'])
        ws.cell(row=row, column=4, value=item['total_nilai'])
        for c in range(1, 5):
            _style_cell(ws, row, c)
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=f'Dicetak: {datetime.now().strftime("%d/%m/%Y %H:%M")}').font = NORMAL_FONT

    _auto_width(ws, 6)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_monthly(bulan: str | None, tahun: str | None) -> bytes:
    """Export monthly report as .xlsx bytes (format SOP)."""
    data = monthly_report(bulan, tahun)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Laporan Bulanan'

    periode = data.get('periode', {})
    periode_str = f'{periode.get("mulai", "")} s.d. {periode.get("selesai", "")}'
    row = _write_kop(ws, f'LAPORAN BULANAN (Bulan {data["bulan"]}, {data["tahun"]})', periode_str)

    # Key metrics table
    metrics = [
        ('Jumlah Nasabah Terdaftar', data['jumlah_nasabah_terdaftar']),
        ('Jumlah Nasabah Aktif', data['jumlah_nasabah_aktif']),
        ('Jumlah Transaksi', data['jumlah_transaksi']),
        ('Total Sampah (kg)', data['total_sampah_kg']),
        ('Total Nilai Setoran', data['total_nilai_setoran']),
        ('Total Penarikan', data['total_penarikan']),
        ('Total Saldo Beredar', data['total_saldo_beredar']),
        ('Jumlah Reward Ditukar', data['jumlah_reward_ditukar']),
    ]
    headers = ['Indikator', 'Nilai']
    ws.cell(row=row, column=1, value='Indikator')
    ws.cell(row=row, column=2, value='Nilai')
    _style_header_row(ws, row, 2)
    row += 1

    for label, value in metrics:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=str(value) if value is not None else '0')
        _style_cell(ws, row, 1, bold=True)
        _style_cell(ws, row, 2)
        row += 1
    row += 1

    # Tonase per jenis
    row = _write_section_title(ws, row, 'Tonase per Jenis')
    tonase_headers = ['No', 'Kategori', 'Total Berat (kg)', 'Total Nilai']
    for i, h in enumerate(tonase_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(tonase_headers))
    row += 1

    for idx, item in enumerate(data.get('tonase_per_jenis', []), 1):
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=item['nama'])
        ws.cell(row=row, column=3, value=item['total_berat_kg'])
        ws.cell(row=row, column=4, value=item['total_nilai'])
        for c in range(1, 5):
            _style_cell(ws, row, c)
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=f'Dicetak: {datetime.now().strftime("%d/%m/%Y %H:%M")}').font = NORMAL_FONT

    _auto_width(ws, 6)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_waste(start: str | None, end: str | None) -> bytes:
    """Export waste report as .xlsx bytes."""
    data = waste_report(start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Laporan Sampah'

    periode = data.get('periode', {})
    periode_str = f'{periode.get("mulai", "")} s.d. {periode.get("selesai", "")}'
    row = _write_kop(ws, 'LAPORAN SAMPAH PER KATEGORI', periode_str)

    # Summary
    headers = ['Total Berat (kg)', 'Total Nilai']
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(headers))
    row += 1

    ws.cell(row=row, column=1, value=data['total_berat_kg'])
    ws.cell(row=row, column=2, value=data['total_nilai'])
    for c in range(1, 3):
        _style_cell(ws, row, c)
    row += 2

    # Per kategori
    row = _write_section_title(ws, row, 'Detail per Kategori')
    kat_headers = ['No', 'Kategori', 'Total Berat (kg)', 'Total Nilai']
    for i, h in enumerate(kat_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(kat_headers))
    row += 1

    for idx, item in enumerate(data.get('per_kategori', []), 1):
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=item['nama'])
        ws.cell(row=row, column=3, value=item['total_berat_kg'])
        ws.cell(row=row, column=4, value=item['total_nilai'])
        for c in range(1, 5):
            _style_cell(ws, row, c)
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=f'Dicetak: {datetime.now().strftime("%d/%m/%Y %H:%M")}').font = NORMAL_FONT

    _auto_width(ws, 6)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_evaluation(start: str | None, end: str | None) -> bytes:
    """Export evaluation report as .xlsx bytes (format SOP)."""
    data = evaluation_report(start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Laporan Evaluasi'

    periode = data.get('periode', {})
    periode_str = f'{periode.get("mulai", "")} s.d. {periode.get("selesai", "")}'
    row = _write_kop(ws, 'LAPORAN EVALUASI PROGRAM', periode_str)

    # Key indicators
    indicators = [
        ('Jumlah Nasabah Terdaftar', data['jumlah_nasabah_terdaftar']),
        ('Jumlah Nasabah Aktif', data['jumlah_nasabah_aktif']),
        ('Nasabah Baru', data['nasabah_baru']),
        ('Jumlah Transaksi', data['jumlah_transaksi']),
        ('Total Sampah (kg)', data['total_sampah_kg']),
        ('Total Nilai Setoran', data['total_nilai_setoran']),
        ('Total Penarikan', data['total_penarikan']),
        ('Jumlah Reward Ditukar', data['jumlah_reward_ditukar']),
        ('Total Poin Ditukar', data['total_poin_ditukar']),
    ]
    headers = ['Indikator', 'Nilai']
    ws.cell(row=row, column=1, value='Indikator')
    ws.cell(row=row, column=2, value='Nilai')
    _style_header_row(ws, row, 2)
    row += 1

    for label, value in indicators:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=str(value) if value is not None else '0')
        _style_cell(ws, row, 1, bold=True)
        _style_cell(ws, row, 2)
        row += 1
    row += 1

    # Wilayah teraktif
    wilayah = data.get('wilayah_teraktif', [])
    if wilayah:
        row = _write_section_title(ws, row, 'Wilayah Teraktif')
        w_headers = ['No', 'Kelurahan', 'Nasabah Aktif']
        for i, h in enumerate(w_headers, 1):
            ws.cell(row=row, column=i, value=h)
        _style_header_row(ws, row, len(w_headers))
        row += 1
        for idx, item in enumerate(wilayah, 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=item['kelurahan'])
            ws.cell(row=row, column=3, value=item['jumlah_nasabah_aktif'])
            for c in range(1, 4):
                _style_cell(ws, row, c)
            row += 1
        row += 1

    # Kendala per jenis
    kendala = data.get('kendala', {})
    if kendala.get('per_jenis'):
        row = _write_section_title(ws, row, f'Kendala Lapangan (Total: {kendala["total_pengaduan"]})')
        k_headers = ['No', 'Jenis Kendala', 'Jumlah']
        for i, h in enumerate(k_headers, 1):
            ws.cell(row=row, column=i, value=h)
        _style_header_row(ws, row, len(k_headers))
        row += 1
        for idx, item in enumerate(kendala['per_jenis'], 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=item['label'])
            ws.cell(row=row, column=3, value=item['jumlah'])
            for c in range(1, 4):
                _style_cell(ws, row, c)
            row += 1
        row += 1

    # Rekomendasi
    rekomendasi = data.get('rekomendasi', [])
    if rekomendasi:
        row = _write_section_title(ws, row, 'Rekomendasi Tindak Lanjut')
        for idx, item in enumerate(rekomendasi, 1):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            ws.cell(row=row, column=1, value=f'{idx}. {item}')
            _style_cell(ws, row, 1)
            row += 1
        row += 1

    # Tonase per jenis
    row = _write_section_title(ws, row, 'Tonase per Jenis')
    tonase_headers = ['No', 'Kategori', 'Total Berat (kg)', 'Total Nilai']
    for i, h in enumerate(tonase_headers, 1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(tonase_headers))
    row += 1

    for idx, item in enumerate(data.get('tonase_per_jenis', []), 1):
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=item['nama'])
        ws.cell(row=row, column=3, value=item['total_berat_kg'])
        ws.cell(row=row, column=4, value=item['total_nilai'])
        for c in range(1, 5):
            _style_cell(ws, row, c)
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=f'Dicetak: {datetime.now().strftime("%d/%m/%Y %H:%M")}').font = NORMAL_FONT

    _auto_width(ws, 6)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
